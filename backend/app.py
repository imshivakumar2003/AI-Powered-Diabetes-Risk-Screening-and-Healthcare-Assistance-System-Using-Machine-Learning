"""
Diabetes Prediction API
------------------------
Flask backend that wraps the existing trained model (diabetes_prediction_model.pkl).

IMPORTANT: The prediction logic and input encoding below are copied over
UNCHANGED from the original Streamlit app.py, so predictions remain identical.

DATABASE: Prediction history is now stored in a SQL database (SQLite by
default, see database.py) instead of a JSON file. All existing API routes,
request formats, and response shapes are unchanged.
"""

import os
import io
import csv
import pickle
import re
import traceback
import sys

import numpy as np
from dotenv import load_dotenv
from flask import Flask, jsonify, request, send_file
from flask_cors import CORS
from sqlalchemy.exc import SQLAlchemyError

load_dotenv()  # load backend/.env automatically on app start

from database import init_engine_tables, get_session
from models import PredictionHistory, Report
from routes.auth_routes import auth_bp
from routes.admin_routes import admin_bp
from routes.report_routes import report_bp
from routes.chat_routes import chat_bp
from routes.recommendations_routes import recommendations_bp
from routes.notifications_routes import notifications_bp
from auth import get_optional_user
from audit import log_action

# ---------------------------------------------------------------------------
# App setup
# ---------------------------------------------------------------------------
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
MODEL_PATH = os.path.join(BASE_DIR, "diabetes_prediction_model.pkl")

from flask import Flask, jsonify, request, send_file, make_response

app = Flask(__name__)
CORS(app, resources={r"/api/*": {"origins": "*", "methods": ["GET", "POST", "PUT", "DELETE", "OPTIONS"], "allow_headers": ["Content-Type", "Authorization"]}})

@app.errorhandler(404)
def not_found_handler(e):
    if request.method == "OPTIONS":
        res = make_response("", 200)
        res.headers["Access-Control-Allow-Origin"] = "*"
        res.headers["Access-Control-Allow-Headers"] = "Content-Type, Authorization, *"
        res.headers["Access-Control-Allow-Methods"] = "GET, POST, PUT, DELETE, OPTIONS"
        return res
    return jsonify({"error": "Resource not found."}), 404

@app.errorhandler(405)
def method_not_allowed_handler(e):
    if request.method == "OPTIONS":
        res = make_response("", 200)
        res.headers["Access-Control-Allow-Origin"] = "*"
        res.headers["Access-Control-Allow-Headers"] = "Content-Type, Authorization, *"
        res.headers["Access-Control-Allow-Methods"] = "GET, POST, PUT, DELETE, OPTIONS"
        return res
    return jsonify({"error": "Method not allowed."}), 405

@app.after_request
def after_request(response):
    if request.method == "OPTIONS":
        response.status_code = 200
    response.headers["Access-Control-Allow-Origin"] = "*"
    response.headers["Access-Control-Allow-Headers"] = "Content-Type, Authorization, *"
    response.headers["Access-Control-Allow-Methods"] = "GET, POST, PUT, DELETE, OPTIONS"
    return response

# New: authentication routes (register/login/forgot-password/profile/etc.)
# mounted at /api/auth/*. This is purely additive — none of the existing
# routes below (/api/predict, /api/history, /api/stats, /api/health) change.
app.register_blueprint(auth_bp)
app.register_blueprint(admin_bp)
app.register_blueprint(report_bp)
app.register_blueprint(chat_bp)
app.register_blueprint(recommendations_bp)
app.register_blueprint(notifications_bp)

@app.route("/api/speech-to-text", methods=["POST", "OPTIONS"])
def direct_speech_to_text():
    if request.method == "OPTIONS":
        return "", 200

    audio_file = request.files.get("audio") or request.files.get("file")
    if not audio_file:
        return jsonify({"error": "No audio file provided."}), 400

    language = request.form.get("language", "en")
    whisper_lang_map = {"en": "en", "hi": "hi", "kn": "kn", "ta": "ta", "te": "te"}
    target_lang = whisper_lang_map.get(language, "en")

    try:
        from groq_service import get_groq_client
        client = get_groq_client()
        if client:
            filename = audio_file.filename or "recording.webm"
            audio_bytes = audio_file.read()
            if len(audio_bytes) == 0:
                return jsonify({"error": "Audio recording is empty."}), 400

            transcription = client.audio.transcriptions.create(
                file=(filename, audio_bytes),
                model="whisper-large-v3",
                language=target_lang,
                response_format="json",
            )
            text = transcription.text.strip() if hasattr(transcription, "text") else str(transcription).strip()
            return jsonify({"text": text, "transcript": text, "engine": "groq-whisper"})
    except Exception as exc:
        import logging
        logging.getLogger(__name__).error(f"Groq Whisper STT failed: {exc}")

    return jsonify({"error": "Groq Whisper API transcription failed. Please retry."}), 500


@app.route("/api/text-to-speech", methods=["POST", "OPTIONS"])
def direct_text_to_speech():
    if request.method == "OPTIONS":
        return "", 200

    payload = request.get_json(silent=True) or {}
    text = payload.get("text", "").strip()
    language = payload.get("language", "en")

    if not text:
        return jsonify({"error": "No text provided."}), 400

    clean_text = re.sub(r'[*#`_-]', '', text)
    lang_map = {"en": "en", "hi": "hi", "kn": "kn", "ta": "ta", "te": "te"}
    target_lang = lang_map.get(language, "en")

    try:
        from gtts import gTTS
        import io

        tts = gTTS(text=clean_text[:1200], lang=target_lang, slow=False)
        fp = io.BytesIO()
        tts.write_to_fp(fp)
        fp.seek(0)
        return send_file(fp, mimetype="audio/mp3", as_attachment=False)
    except Exception as exc:
        import logging
        logging.getLogger(__name__).error(f"TTS synthesis error: {exc}")
        return jsonify({"error": "Text-to-Speech synthesis failed."}), 500

# Create DB tables automatically if they don't exist yet. Safe to call on
# every startup — never touches existing tables or data.
init_engine_tables()

# ---------------------------------------------------------------------------
# Load model once at startup
# ---------------------------------------------------------------------------
with open(MODEL_PATH, "rb") as f:
    loaded_model = pickle.load(f)

# ---------------------------------------------------------------------------
# Encoding maps (MUST match training — copied from original app.py)
# ---------------------------------------------------------------------------
GENDER_MAP = {"Male": 0, "Female": 1}
SMOKING_MAP = {
    "never": 0,
    "former": 1,
    "current": 2,
    "not current": 3,
    "ever": 4,
}

REQUIRED_FIELDS = [
    "age",
    "hypertension",
    "heart_disease",
    "bmi",
    "HbA1c_level",
    "blood_glucose_level",
    "gender",
    "smoking_history",
]


def validate_payload(payload):
    missing = [field for field in REQUIRED_FIELDS if field not in payload]
    if missing:
        return f"Missing fields: {', '.join(missing)}"
    if payload["gender"] not in GENDER_MAP:
        return "Invalid value for 'gender'. Expected 'Male' or 'Female'."
    if payload["smoking_history"] not in SMOKING_MAP:
        return f"Invalid value for 'smoking_history'. Expected one of {list(SMOKING_MAP)}."
    return None


# ---------------------------------------------------------------------------
# Routes
# ---------------------------------------------------------------------------
@app.route("/api/health", methods=["GET"])
def health():
    db_ok = True
    try:
        with get_session() as db:
            db.query(PredictionHistory).limit(1).all()
    except SQLAlchemyError:
        db_ok = False

    return jsonify({
        "status": "ok" if db_ok else "degraded",
        "model_loaded": loaded_model is not None,
        "database_connected": db_ok,
    })


@app.route("/api/predict", methods=["POST"])
def predict():
    raw_body = request.get_data(as_text=True)
    payload = request.get_json(silent=True) or {}
    print(f"\n--- [PREDICT API REQUEST RECEIVED] ---", flush=True)
    print(f"Raw Body: {raw_body}", flush=True)
    print(f"Parsed Payload: {payload}", flush=True)

    try:
        error = validate_payload(payload)
        if error:
            print(f"[PREDICT VALIDATION FAILED] {error}", flush=True)
            return jsonify({"success": False, "error": error}), 400

        gender_encoded = GENDER_MAP[payload["gender"]]
        smoking_encoded = SMOKING_MAP[payload["smoking_history"]]

        hypertension_bool = str(payload["hypertension"]) in ("Yes", "1", "true", "True")
        heart_disease_bool = str(payload["heart_disease"]) in ("Yes", "1", "true", "True")

        user_input = np.array([[
            float(payload["age"]),
            1 if hypertension_bool else 0,
            1 if heart_disease_bool else 0,
            float(payload["bmi"]),
            float(payload["HbA1c_level"]),
            float(payload["blood_glucose_level"]),
            gender_encoded,
            smoking_encoded,
        ]])

        if loaded_model is None:
            raise RuntimeError("Scikit-learn ML model (diabetes_prediction_model.pkl) is not loaded.")

        prediction = loaded_model.predict(user_input)
        result = "Positive" if int(prediction[0]) == 1 else "Negative"

        probability = None
        if hasattr(loaded_model, "predict_proba"):
            proba = loaded_model.predict_proba(user_input)[0]
            probability = round(float(proba[1]) * 100, 2)

        print(f"Prediction Output: Result={result}, Probability={probability}%", flush=True)

        current_user = get_optional_user()
        user_id = current_user["id"] if current_user else None
        print(f"Database Insert Payload: user_id={user_id}, age={payload['age']}, result={result}, prob={probability}", flush=True)

        with get_session() as db:
            row = PredictionHistory(
                user_id=user_id,
                age=int(payload["age"]),
                gender=payload["gender"],
                hypertension=hypertension_bool,
                heart_disease=heart_disease_bool,
                bmi=float(payload["bmi"]),
                hba1c_level=float(payload["HbA1c_level"]),
                blood_glucose_level=float(payload["blood_glucose_level"]),
                smoking_history=payload["smoking_history"],
                result=result,
                probability=probability,
            )
            db.add(row)
            db.flush()

            existing_report = db.query(Report).filter_by(prediction_id=row.id).first()
            if not existing_report:
                report_row = Report(
                    prediction_id=row.id,
                    user_id=user_id,
                    report_type="pdf",
                )
                db.add(report_row)
                db.flush()
            record = row.to_api_dict()

        log_action(user_id, "prediction_created", f"Result: {result}")
        if user_id:
            try:
                from notification_service import create_user_notification
                create_user_notification(user_id, "prediction_completed")
                create_user_notification(user_id, "report_generated")
                create_user_notification(user_id, "ai_recommendation")
                if result == "Positive" or (probability and probability >= 40.0):
                    create_user_notification(user_id, "high_risk_alert")
            except Exception as e:
                print(f"Notification creation log: {e}", flush=True)

        print("--- [PREDICT API SUCCESS] ---\n", flush=True)
        return jsonify(record)

    except (ValueError, TypeError) as exc:
        tb = traceback.format_exc()
        print(f"\n!!! [PREDICT VALUE/TYPE ERROR] !!!\n{tb}\n", flush=True)
        return jsonify({"success": False, "error": f"Invalid input values: {exc}"}), 400

    except Exception as exc:
        tb = traceback.format_exc()
        print(f"\n!!! [PREDICT API EXCEPTION] !!!", flush=True)
        print(f"Error: {exc}", flush=True)
        print(f"Traceback:\n{tb}", flush=True)
        print("-------------------------------\n", flush=True)
        return jsonify({
            "success": False,
            "error": str(exc),
            "details": f"{type(exc).__name__}: {exc}"
        }), 500


@app.route("/api/history", methods=["GET"])
def get_history():
    limit = request.args.get("limit", default=50, type=int)
    current_user = get_optional_user()

    try:
        with get_session() as db:
            query = db.query(PredictionHistory)
            # If logged in, patients only see their own history. If not
            # logged in (existing frontend behavior, pre-auth), nothing
            # changes — all records are returned, same as before.
            if current_user:
                query = query.filter(PredictionHistory.user_id == current_user["id"])
            rows = query.order_by(PredictionHistory.created_at.desc()).limit(limit).all()
            records = [r.to_api_dict() for r in rows]
    except SQLAlchemyError as exc:
        return jsonify({"error": f"Database error while fetching history: {exc}"}), 500

    return jsonify(records)


def _build_filtered_history_query(db, current_user):
    """Shared filter logic for /api/history/search and /api/history/export."""
    query = db.query(PredictionHistory)
    if current_user:
        query = query.filter(PredictionHistory.user_id == current_user["id"])

    result = request.args.get("result")
    if result in ("Positive", "Negative"):
        query = query.filter(PredictionHistory.result == result)

    gender = request.args.get("gender")
    if gender in ("Male", "Female"):
        query = query.filter(PredictionHistory.gender == gender)

    date_from = request.args.get("date_from")
    if date_from:
        query = query.filter(PredictionHistory.created_at >= date_from)

    date_to = request.args.get("date_to")
    if date_to:
        query = query.filter(PredictionHistory.created_at <= date_to)

    min_age = request.args.get("min_age", type=int)
    if min_age is not None:
        query = query.filter(PredictionHistory.age >= min_age)

    max_age = request.args.get("max_age", type=int)
    if max_age is not None:
        query = query.filter(PredictionHistory.age <= max_age)

    return query


@app.route("/api/history/search", methods=["GET"])
def search_history():
    """Paginated, filterable history search — additive, does not change the
    existing /api/history contract used elsewhere in the app.

    Query params: page, page_size, result, gender, date_from, date_to,
    min_age, max_age.
    """
    page = max(request.args.get("page", default=1, type=int), 1)
    page_size = min(request.args.get("page_size", default=10, type=int), 100)
    current_user = get_optional_user()

    try:
        with get_session() as db:
            query = _build_filtered_history_query(db, current_user)
            total = query.count()
            rows = (
                query.order_by(PredictionHistory.created_at.desc())
                .offset((page - 1) * page_size)
                .limit(page_size)
                .all()
            )
            items = [r.to_api_dict() for r in rows]
    except SQLAlchemyError as exc:
        return jsonify({"error": f"Database error while searching history: {exc}"}), 500

    return jsonify({
        "items": items,
        "total": total,
        "page": page,
        "page_size": page_size,
        "total_pages": (total + page_size - 1) // page_size if page_size else 0,
    })


@app.route("/api/history/export", methods=["GET"])
def export_history():
    """Export filtered history as CSV, Excel, or PDF.
    Query params: format=csv|excel|pdf, plus the same filters as /search.
    """
    fmt = request.args.get("format", default="csv").lower()
    current_user = get_optional_user()

    try:
        with get_session() as db:
            query = _build_filtered_history_query(db, current_user)
            rows = query.order_by(PredictionHistory.created_at.desc()).all()
            records = [r.to_api_dict() for r in rows]
    except SQLAlchemyError as exc:
        return jsonify({"error": f"Database error while exporting history: {exc}"}), 500

    columns = ["id", "timestamp", "age", "gender", "bmi", "HbA1c_level",
               "blood_glucose_level", "hypertension", "heart_disease",
               "smoking_history", "result", "probability"]

    def flat_row(r):
        return [
            r["id"], r["timestamp"], r["input"]["age"], r["input"]["gender"],
            r["input"]["bmi"], r["input"]["HbA1c_level"], r["input"]["blood_glucose_level"],
            r["input"]["hypertension"], r["input"]["heart_disease"],
            r["input"]["smoking_history"], r["result"], r["probability"],
        ]

    if fmt == "csv":
        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow(columns)
        for r in records:
            writer.writerow(flat_row(r))
        mem = io.BytesIO(buf.getvalue().encode("utf-8"))
        return send_file(mem, mimetype="text/csv", as_attachment=True,
                          download_name="prediction_history.csv")

    if fmt == "excel":
        try:
            import openpyxl
            from openpyxl.utils import get_column_letter
        except ImportError:
            return jsonify({"error": "openpyxl is not installed. Run: pip install openpyxl"}), 500

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Prediction History"
        ws.append(columns)
        for r in records:
            ws.append(flat_row(r))
        for i, col in enumerate(columns, start=1):
            ws.column_dimensions[get_column_letter(i)].width = max(12, len(col) + 2)

        mem = io.BytesIO()
        wb.save(mem)
        mem.seek(0)
        return send_file(
            mem,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name="prediction_history.xlsx",
        )

    if fmt == "pdf":
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import landscape, A4
            from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
            from reportlab.lib.styles import getSampleStyleSheet
        except ImportError:
            return jsonify({"error": "reportlab is not installed. Run: pip install reportlab"}), 500

        mem = io.BytesIO()
        doc = SimpleDocTemplate(mem, pagesize=landscape(A4))
        styles = getSampleStyleSheet()
        elements = [Paragraph("Diabetes Screening History", styles["Title"]), Spacer(1, 12)]

        table_data = [columns] + [[str(v) for v in flat_row(r)] for r in records]
        table = Table(table_data, repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f1b2d")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTSIZE", (0, 0), (-1, -1), 7),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#e3e8ee")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f6f8fa")]),
        ]))
        elements.append(table)
        doc.build(elements)
        mem.seek(0)
        return send_file(mem, mimetype="application/pdf", as_attachment=True,
                          download_name="prediction_history.pdf")

    return jsonify({"error": "format must be one of: csv, excel, pdf"}), 400


@app.route("/api/history/<record_id>", methods=["DELETE"])
def delete_history_item(record_id):
    try:
        with get_session() as db:
            row = db.query(PredictionHistory).filter_by(id=record_id).first()
            if row is None:
                return jsonify({"error": "Record not found"}), 404
            db.delete(row)
    except SQLAlchemyError as exc:
        return jsonify({"error": f"Database error while deleting record: {exc}"}), 500

    current_user = get_optional_user()
    log_action(current_user["id"] if current_user else None, "history_deleted", f"Record {record_id}")
    return jsonify({"deleted": record_id})


@app.route("/api/history", methods=["DELETE"])
def clear_history():
    try:
        with get_session() as db:
            db.query(PredictionHistory).delete()
    except SQLAlchemyError as exc:
        return jsonify({"error": f"Database error while clearing history: {exc}"}), 500

    current_user = get_optional_user()
    log_action(current_user["id"] if current_user else None, "history_cleared", "All history cleared")
    return jsonify({"cleared": True})


@app.route("/api/stats", methods=["GET"])
def stats():
    try:
        with get_session() as db:
            rows = db.query(PredictionHistory).all()
            # Extract everything needed while the session is still open —
            # ORM objects become "detached" once the session closes, and
            # touching their attributes afterward raises DetachedInstanceError.
            total = len(rows)
            positive = sum(1 for r in rows if r.result == "Positive")
            avg_glucose = round(sum(r.blood_glucose_level for r in rows) / total, 1) if total else 0
            avg_bmi = round(sum(r.bmi for r in rows) / total, 1) if total else 0
    except SQLAlchemyError as exc:
        return jsonify({"error": f"Database error while computing stats: {exc}"}), 500

    negative = total - positive

    return jsonify({
        "total_predictions": total,
        "positive": positive,
        "negative": negative,
        "avg_glucose": avg_glucose,
        "avg_bmi": avg_bmi,
    })


if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    debug = os.environ.get("FLASK_DEBUG", "1") == "1"
    # use_reloader is disabled to avoid the reloader restarting the process
    # mid-request (also prevents interference with the SQLite file being
    # opened/written during predictions).
    app.run(debug=debug, port=port, use_reloader=False)
